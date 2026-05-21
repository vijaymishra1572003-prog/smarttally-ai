import os
import csv
import json
import logging

logger = logging.getLogger(__name__)


class OCREngine:
    """Multi-engine OCR processor supporting PDF, images, CSV, and Excel files."""

    def extract_text(self, file_path, file_type):
        file_type = file_type.lower()
        if file_type in ("csv",):
            return self._extract_csv(file_path)
        if file_type in ("xlsx", "xls"):
            return self._extract_excel(file_path)
        if file_type == "pdf":
            return self._extract_pdf(file_path)
        if file_type in ("png", "jpg", "jpeg"):
            return self._extract_image(file_path)
        raise ValueError(f"Unsupported file type: {file_type}")

    def _extract_pdf(self, file_path):
        text = ""
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            if text.strip():
                return text.strip()
        except ImportError:
            logger.warning("pdfplumber not installed, trying PyPDF2")
        except Exception as e:
            logger.warning(f"pdfplumber failed: {e}")

        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            if text.strip():
                return text.strip()
        except ImportError:
            logger.warning("PyPDF2 not installed")
        except Exception as e:
            logger.warning(f"PyPDF2 failed: {e}")

        return self._ocr_fallback(file_path)

    def _extract_image(self, file_path):
        try:
            from paddleocr import PaddleOCR
            ocr = PaddleOCR(use_angle_cls=True, lang="en", show_log=False)
            result = ocr.ocr(file_path, cls=True)
            lines = []
            if result and result[0]:
                for line in result[0]:
                    if line[1]:
                        lines.append(line[1][0])
            if lines:
                return "\n".join(lines)
        except ImportError:
            logger.warning("PaddleOCR not available")
        except Exception as e:
            logger.warning(f"PaddleOCR failed: {e}")

        try:
            import pytesseract
            from PIL import Image
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img)
            if text.strip():
                return text.strip()
        except ImportError:
            logger.warning("pytesseract not installed")
        except Exception as e:
            logger.warning(f"Tesseract failed: {e}")

        return f"[OCR extraction attempted for {os.path.basename(file_path)}]"

    def _ocr_fallback(self, file_path):
        return self._extract_image(file_path)

    def _extract_csv(self, file_path):
        rows = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(" | ".join(row))
        return "\n".join(rows)

    def _extract_excel(self, file_path):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    lines.append(" | ".join(cells))
            return "\n".join(lines)
        except ImportError:
            logger.warning("openpyxl not installed")
            return "[Excel extraction requires openpyxl]"
